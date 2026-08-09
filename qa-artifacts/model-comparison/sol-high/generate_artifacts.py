from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


SOURCE = Path(
    "/Users/jyb-m3max/Desktop/codex/genoffice/qa-artifacts/"
    "gpt-5.6-luna-xhigh-smoke/gpt-5.6-luna-xhigh-saas-portfolio.xlsx"
)
OUT_DIR = Path(
    "/Users/jyb-m3max/Desktop/codex/genoffice/qa-artifacts/"
    "model-comparison/sol-high"
)
XLSX_PATH = OUT_DIR / "sol-high-portfolio.xlsx"
PPTX_PATH = OUT_DIR / "sol-high-investment-committee.pptx"
METRICS_PATH = OUT_DIR / "computed-metrics.json"

NAVY = "0B132B"
NAVY_2 = "14213D"
NAVY_3 = "1E2A49"
MINT = "5EEAD4"
MINT_DARK = "0F766E"
WHITE = "F8FAFC"
PALE = "E2E8F0"
SLATE = "64748B"
GREEN = "DCFCE7"
RED = "FEE2E2"
AMBER = "FEF3C7"
RED_TEXT = "B91C1C"
GREEN_TEXT = "166534"
FONT_KR = "Nanum Gothic"
THIN = Side(style="thin", color="CBD5E1")


def extract_source() -> tuple[list[dict], dict]:
    wb = load_workbook(SOURCE, data_only=False, read_only=False)
    assert wb.sheetnames == ["Portfolio", "Scenarios", "Dashboard"]
    ws = wb["Portfolio"]
    weights = {
        "strategy": float(ws["B4"].value),
        "finance": float(ws["C4"].value),
        "risk": float(ws["D4"].value),
    }
    items: list[dict] = []
    for row in range(8, 18):
        item = {
            "name": ws.cell(row, 1).value,
            "arr": int(ws.cell(row, 2).value),
            "prob": float(ws.cell(row, 3).value),
            "dev": int(ws.cell(row, 4).value),
            "monthly": int(ws.cell(row, 5).value),
            "strategy": int(ws.cell(row, 6).value),
            "risk": int(ws.cell(row, 7).value),
        }
        item["expected_arr"] = item["arr"] * item["prob"]
        item["cost_12m"] = item["dev"] + 12 * item["monthly"]
        item["base_net"] = item["expected_arr"] - item["cost_12m"]
        item["roi"] = item["base_net"] / item["cost_12m"]
        item["score"] = (
            weights["strategy"] * item["strategy"]
            + weights["finance"] * item["roi"] * 100
            + weights["risk"] * (100 - item["risk"])
        )
        items.append(item)

    scenarios_ws = wb["Scenarios"]
    scenarios = []
    for row in range(4, 7):
        scenarios.append(
            {
                "name": scenarios_ws.cell(row, 1).value,
                "arr_multiplier": float(scenarios_ws.cell(row, 2).value),
                "probability_adjustment": float(scenarios_ws.cell(row, 3).value),
                "cost_multiplier": float(scenarios_ws.cell(row, 4).value),
            }
        )
    return items, {"weights": weights, "scenarios": scenarios}


def compute_metrics(items: list[dict], assumptions: dict) -> dict:
    ranked = sorted(items, key=lambda item: item["score"], reverse=True)
    for rank, item in enumerate(ranked, start=1):
        item["rank"] = rank

    scenario_totals: dict[str, float] = {}
    scenario_values: dict[str, list[float]] = {}
    for scenario in assumptions["scenarios"]:
        values = []
        for item in items:
            values.append(
                item["arr"]
                * scenario["arr_multiplier"]
                * (item["prob"] + scenario["probability_adjustment"])
                - item["cost_12m"] * scenario["cost_multiplier"]
            )
        scenario_values[scenario["name"]] = values
        scenario_totals[scenario["name"]] = sum(values)

    total_arr = sum(item["arr"] for item in items)
    top3 = ranked[:3]
    return {
        "investment_count": len(items),
        "total_arr": total_arr,
        "total_expected_arr": sum(item["expected_arr"] for item in items),
        "total_investment": sum(item["cost_12m"] for item in items),
        "base_net_value": sum(item["base_net"] for item in items),
        "weighted_success_probability": sum(
            item["arr"] * item["prob"] for item in items
        )
        / total_arr,
        "weighted_attrition_probability": 1
        - sum(item["arr"] * item["prob"] for item in items) / total_arr,
        "scenario_totals": scenario_totals,
        "scenario_values": scenario_values,
        "ranked": ranked,
        "top3": top3,
        "top3_total_investment": sum(item["cost_12m"] for item in top3),
        "top3_expected_arr": sum(item["expected_arr"] for item in top3),
        "top3_base_net": sum(item["base_net"] for item in top3),
    }


def style_title(ws, cell_range: str, title: str) -> None:
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = title
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name=FONT_KR, size=20, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[cell.row].height = 34


def apply_header(row_cells) -> None:
    for cell in row_cells:
        cell.fill = PatternFill("solid", fgColor=NAVY_2)
        cell.font = Font(name=FONT_KR, size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=MINT))


def set_widths(ws, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def build_workbook(items: list[dict], assumptions: dict, metrics: dict) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.properties.creator = "OpenAI Codex"
    wb.properties.title = "Sol High B2B SaaS Portfolio"
    wb.properties.subject = "GenOffice model comparison investment portfolio"
    wb.properties.keywords = "gpt-5.6-sol, high, portfolio, investment committee"

    portfolio = wb.create_sheet("Portfolio")
    scenarios = wb.create_sheet("Scenarios")
    dashboard = wb.create_sheet("Dashboard")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_setup.orientation = "landscape"
        ws.page_margins.left = 0.2
        ws.page_margins.right = 0.2
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35

    # Portfolio
    style_title(portfolio, "A1:L1", "B2B SaaS 이니셔티브 포트폴리오 평가")
    portfolio["A3"] = "가정"
    for col, label in enumerate(
        ["전략 가중치", "재무 가중치", "리스크 가중치", "가중치 합계", "100% 검증"],
        start=2,
    ):
        portfolio.cell(3, col).value = label
    portfolio["A4"] = "기준"
    portfolio["B4"] = assumptions["weights"]["strategy"]
    portfolio["C4"] = assumptions["weights"]["finance"]
    portfolio["D4"] = assumptions["weights"]["risk"]
    portfolio["E4"] = "=SUM(B4:D4)"
    portfolio["F4"] = '=IF(E4=1,"검증 완료","오류")'
    apply_header(portfolio[3][0:6])
    for cell in portfolio[4][0:6]:
        cell.fill = PatternFill("solid", fgColor="F1F5F9")
        cell.font = Font(name=FONT_KR, size=10, color=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN)
    for cell in portfolio["B4:D4"]:
        pass
    for addr in ["B4", "C4", "D4", "E4"]:
        portfolio[addr].number_format = "0%"

    headers = [
        "이니셔티브",
        "예상 ARR(원)",
        "성공확률",
        "초기개발비(원)",
        "월운영비(원)",
        "전략점수",
        "리스크점수",
        "기대 ARR(원)",
        "12개월 순가치(원)",
        "ROI",
        "리스크조정점수",
        "우선순위",
    ]
    for col, header in enumerate(headers, start=1):
        portfolio.cell(7, col).value = header
    apply_header(portfolio[7])
    portfolio.row_dimensions[7].height = 34
    for index, item in enumerate(items, start=8):
        values = [
            item["name"],
            item["arr"],
            item["prob"],
            item["dev"],
            item["monthly"],
            item["strategy"],
            item["risk"],
        ]
        for col, value in enumerate(values, start=1):
            portfolio.cell(index, col).value = value
        portfolio.cell(index, 8).value = f"=B{index}*C{index}"
        portfolio.cell(index, 9).value = f"=H{index}-(D{index}+12*E{index})"
        portfolio.cell(index, 10).value = f"=I{index}/(D{index}+12*E{index})"
        portfolio.cell(index, 11).value = (
            f"=$B$4*F{index}+$C$4*(J{index}*100)+$D$4*(100-G{index})"
        )
        portfolio.cell(index, 12).value = f"=RANK(K{index},$K$8:$K$17,0)"
        fill = "FFFFFF" if index % 2 == 0 else "F8FAFC"
        for cell in portfolio[index]:
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(name=FONT_KR, size=9, color=NAVY)
            cell.alignment = Alignment(
                horizontal="left" if cell.column == 1 else "right",
                vertical="center",
                wrap_text=cell.column == 1,
            )
            cell.border = Border(bottom=THIN)
        portfolio.row_dimensions[index].height = 28
    portfolio["A18"] = "합계 / 가중 평균"
    portfolio["B18"] = "=SUM(B8:B17)"
    portfolio["C18"] = "=SUMPRODUCT(B8:B17,C8:C17)/SUM(B8:B17)"
    portfolio["D18"] = "=SUM(D8:D17)"
    portfolio["E18"] = "=SUM(E8:E17)"
    portfolio["H18"] = "=SUM(H8:H17)"
    portfolio["I18"] = "=SUM(I8:I17)"
    portfolio["J18"] = "=I18/(D18+12*E18)"
    for cell in portfolio[18]:
        cell.fill = PatternFill("solid", fgColor="CCFBF1")
        cell.font = Font(name=FONT_KR, size=10, bold=True, color=NAVY)
        cell.border = Border(top=Side(style="medium", color=MINT_DARK))
    for row in range(8, 19):
        for col in [2, 4, 5, 8, 9]:
            portfolio.cell(row, col).number_format = '₩#,##0;[Red]-₩#,##0'
        for col in [3, 10]:
            portfolio.cell(row, col).number_format = "0.0%"
        for col in [6, 7, 11]:
            portfolio.cell(row, col).number_format = "0.0"
        portfolio.cell(row, 12).number_format = "0"
    portfolio.conditional_formatting.add(
        "G8:G17",
        ColorScaleRule(
            start_type="min", start_color="DCFCE7",
            mid_type="percentile", mid_value=50, mid_color="FEF3C7",
            end_type="max", end_color="FECACA",
        ),
    )
    portfolio.conditional_formatting.add(
        "I8:I17", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED))
    )
    portfolio.conditional_formatting.add(
        "K8:K17",
        ColorScaleRule(
            start_type="min", start_color="FEE2E2",
            mid_type="percentile", mid_value=50, mid_color="FEF3C7",
            end_type="max", end_color="A7F3D0",
        ),
    )
    portfolio.freeze_panes = "A8"
    portfolio.auto_filter.ref = "A7:L17"
    portfolio.print_title_rows = "1:7"
    portfolio.print_area = "A1:L18"
    set_widths(
        portfolio,
        {"A": 34, "B": 17, "C": 12, "D": 17, "E": 16, "F": 11,
         "G": 11, "H": 17, "I": 18, "J": 11, "K": 16, "L": 11},
    )

    # Scenarios
    style_title(scenarios, "A1:D1", "B2B SaaS 시나리오 분석")
    scenario_headers = ["시나리오", "ARR 배수", "성공확률 조정", "비용 배수"]
    for col, header in enumerate(scenario_headers, start=1):
        scenarios.cell(3, col).value = header
    apply_header(scenarios[3][0:4])
    for row, scenario in enumerate(assumptions["scenarios"], start=4):
        scenarios.cell(row, 1).value = scenario["name"]
        scenarios.cell(row, 2).value = scenario["arr_multiplier"]
        scenarios.cell(row, 3).value = scenario["probability_adjustment"]
        scenarios.cell(row, 4).value = scenario["cost_multiplier"]
        for cell in scenarios[row][0:4]:
            cell.fill = PatternFill("solid", fgColor="F8FAFC")
            cell.font = Font(name=FONT_KR, size=10, color=NAVY)
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(bottom=THIN)
        scenarios.cell(row, 2).number_format = "0%"
        scenarios.cell(row, 3).number_format = "+0%;-0%;0%"
        scenarios.cell(row, 4).number_format = "0%"
    detailed_headers = [
        "이니셔티브",
        "Downside 기대 순가치(원)",
        "Base 기대 순가치(원)",
        "Upside 기대 순가치(원)",
    ]
    for col, header in enumerate(detailed_headers, start=1):
        scenarios.cell(9, col).value = header
    apply_header(scenarios[9][0:4])
    scenarios.row_dimensions[9].height = 34
    for source_row, row in zip(range(8, 18), range(10, 20)):
        scenarios.cell(row, 1).value = f"=Portfolio!$A{source_row}"
        for col, assumption_row in zip(range(2, 5), range(4, 7)):
            scenarios.cell(row, col).value = (
                f"=Portfolio!$B{source_row}*$B${assumption_row}*"
                f"(Portfolio!$C{source_row}+$C${assumption_row})-"
                f"(Portfolio!$D{source_row}+12*Portfolio!$E{source_row})*$D${assumption_row}"
            )
        for cell in scenarios[row][0:4]:
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if row % 2 == 0 else "F8FAFC")
            cell.font = Font(name=FONT_KR, size=9, color=NAVY)
            cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "right", wrap_text=True)
            cell.border = Border(bottom=THIN)
        for col in range(2, 5):
            scenarios.cell(row, col).number_format = '₩#,##0;[Red]-₩#,##0'
        scenarios.row_dimensions[row].height = 27
    scenarios["A20"] = "포트폴리오 합계"
    scenarios["B20"] = "=SUM(B10:B19)"
    scenarios["C20"] = "=SUM(C10:C19)"
    scenarios["D20"] = "=SUM(D10:D19)"
    for cell in scenarios[20][0:4]:
        cell.fill = PatternFill("solid", fgColor="CCFBF1")
        cell.font = Font(name=FONT_KR, size=10, bold=True, color=NAVY)
        cell.border = Border(top=Side(style="medium", color=MINT_DARK))
    for col in range(2, 5):
        scenarios.cell(20, col).number_format = '₩#,##0;[Red]-₩#,##0'
    scenarios.conditional_formatting.add(
        "B10:D20",
        ColorScaleRule(
            start_type="min", start_color="FECACA",
            mid_type="num", mid_value=0, mid_color="FFF7ED",
            end_type="max", end_color="A7F3D0",
        ),
    )
    scenarios.freeze_panes = "A10"
    scenarios.auto_filter.ref = "A9:D19"
    scenarios.print_title_rows = "1:9"
    scenarios.print_area = "A1:D20"
    set_widths(scenarios, {"A": 38, "B": 25, "C": 25, "D": 25})

    # Dashboard
    style_title(dashboard, "A1:L1", "B2B SaaS 투자 포트폴리오 대시보드")
    kpis = [
        ("A3:B3", "A4:B4", "총 투자액", "=Portfolio!D18+12*Portfolio!E18", '₩#,##0'),
        ("C3:D3", "C4:D4", "예상 ARR", "=Portfolio!H18", '₩#,##0'),
        ("E3:F3", "E4:F4", "Base 순가치", "=Scenarios!C20", '₩#,##0;[Red]-₩#,##0'),
        ("G3:H3", "G4:H4", "가중 성공확률", "=Portfolio!C18", "0.0%"),
        ("I3:J3", "I4:J4", "가중 이탈률", "=1-G4", "0.0%"),
    ]
    for label_range, value_range, label, formula, num_format in kpis:
        dashboard.merge_cells(label_range)
        dashboard.merge_cells(value_range)
        label_cell = dashboard[label_range.split(":")[0]]
        value_cell = dashboard[value_range.split(":")[0]]
        label_cell.value = label
        value_cell.value = formula
        label_cell.fill = PatternFill("solid", fgColor=NAVY_2)
        value_cell.fill = PatternFill("solid", fgColor="ECFDF5")
        label_cell.font = Font(name=FONT_KR, size=9, bold=True, color=PALE)
        value_cell.font = Font(name=FONT_KR, size=16, bold=True, color=MINT_DARK)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.number_format = num_format
    dashboard.row_dimensions[3].height = 23
    dashboard.row_dimensions[4].height = 34

    dashboard["A7"] = "순위"
    dashboard["B7"] = "우선 투자안"
    dashboard["C7"] = "리스크조정점수"
    dashboard["D7"] = "Base 순가치(원)"
    dashboard["G7"] = "시나리오"
    dashboard["H7"] = "포트폴리오 순가치(원)"
    dashboard["I7"] = "기준 대비(원)"
    dashboard["J7"] = "판정"
    apply_header(dashboard[7][0:4])
    apply_header(dashboard[7][6:10])
    for row in range(8, 13):
        dashboard.cell(row, 1).value = row - 7
        dashboard.cell(row, 2).value = f"=INDEX(Portfolio!$A$8:$A$17,MATCH($A{row},Portfolio!$L$8:$L$17,0))"
        dashboard.cell(row, 3).value = f"=INDEX(Portfolio!$K$8:$K$17,MATCH($A{row},Portfolio!$L$8:$L$17,0))"
        dashboard.cell(row, 4).value = f"=INDEX(Portfolio!$I$8:$I$17,MATCH($A{row},Portfolio!$L$8:$L$17,0))"
        for cell in dashboard[row][0:4]:
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if row % 2 == 0 else "F8FAFC")
            cell.font = Font(name=FONT_KR, size=9, color=NAVY)
            cell.alignment = Alignment(horizontal="left" if cell.column == 2 else "right", vertical="center", wrap_text=True)
            cell.border = Border(bottom=THIN)
        dashboard.cell(row, 3).number_format = "0.0"
        dashboard.cell(row, 4).number_format = '₩#,##0;[Red]-₩#,##0'
        dashboard.row_dimensions[row].height = 27
    for row, source_col in zip(range(8, 11), [2, 3, 4]):
        dashboard.cell(row, 7).value = f"=Scenarios!A{row - 4}"
        dashboard.cell(row, 8).value = f"=Scenarios!{get_column_letter(source_col)}20"
        dashboard.cell(row, 9).value = f"=H{row}-$H$9"
        dashboard.cell(row, 10).value = f'=IF(H{row}<0,"위험","양호")'
        for cell in dashboard[row][6:10]:
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if row % 2 == 0 else "F8FAFC")
            cell.font = Font(name=FONT_KR, size=9, color=NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=THIN)
        for col in [8, 9]:
            dashboard.cell(row, col).number_format = '₩#,##0;[Red]-₩#,##0'

    dashboard.merge_cells("A31:L31")
    dashboard["A31"] = "경영진 권고"
    dashboard["A31"].fill = PatternFill("solid", fgColor=NAVY_2)
    dashboard["A31"].font = Font(name=FONT_KR, size=11, bold=True, color=MINT)
    dashboard["A31"].alignment = Alignment(horizontal="left")
    dashboard.merge_cells("A32:L34")
    dashboard["A32"] = (
        '=IF(Scenarios!B20<0,"Downside에서 순가치가 음수이므로 전액 동시 집행보다 1~3위 투자안의 단계적 집행을 권고합니다.",'
        '"Downside에서도 순가치가 양수이므로 계획대로 집행 가능합니다.")'
    )
    dashboard["A32"].fill = PatternFill("solid", fgColor="ECFDF5")
    dashboard["A32"].font = Font(name=FONT_KR, size=12, bold=True, color=NAVY)
    dashboard["A32"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    dashboard.row_dimensions[32].height = 46

    score_chart = BarChart()
    score_chart.type = "bar"
    score_chart.style = 10
    score_chart.title = "우선순위 Top 5 — 리스크조정점수"
    score_chart.y_axis.title = "투자안"
    score_chart.x_axis.title = "점수"
    score_chart.height = 7.0
    score_chart.width = 13.0
    score_chart.legend = None
    score_chart.add_data(Reference(dashboard, min_col=3, min_row=7, max_row=12), titles_from_data=True)
    score_chart.set_categories(Reference(dashboard, min_col=2, min_row=8, max_row=12))
    score_chart.series[0].graphicalProperties.solidFill = MINT_DARK
    score_chart.series[0].graphicalProperties.line.solidFill = MINT_DARK
    dashboard.add_chart(score_chart, "A15")

    scenario_chart = BarChart()
    scenario_chart.type = "col"
    scenario_chart.style = 10
    scenario_chart.title = "시나리오 스트레스 테스트"
    scenario_chart.y_axis.title = "순가치(원)"
    scenario_chart.height = 7.0
    scenario_chart.width = 13.0
    scenario_chart.legend = None
    scenario_chart.add_data(Reference(dashboard, min_col=8, min_row=7, max_row=10), titles_from_data=True)
    scenario_chart.set_categories(Reference(dashboard, min_col=7, min_row=8, max_row=10))
    scenario_chart.series[0].graphicalProperties.solidFill = MINT
    scenario_chart.series[0].graphicalProperties.line.solidFill = MINT_DARK
    dashboard.add_chart(scenario_chart, "G15")

    dashboard.conditional_formatting.add("C8:C12", DataBarRule(start_type="min", end_type="max", color=MINT_DARK))
    dashboard.conditional_formatting.add(
        "H8:H10", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED))
    )
    dashboard.conditional_formatting.add(
        "H8:H10", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=PatternFill("solid", fgColor=GREEN))
    )
    dashboard.freeze_panes = "A8"
    dashboard.print_area = "A1:L34"
    dashboard.sheet_properties.pageSetUpPr.fitToPage = True
    dashboard.page_setup.fitToWidth = 1
    dashboard.page_setup.fitToHeight = 1
    set_widths(
        dashboard,
        {"A": 9, "B": 27, "C": 14, "D": 17, "E": 12, "F": 12,
         "G": 14, "H": 18, "I": 15, "J": 12, "K": 2, "L": 2},
    )

    wb.save(XLSX_PATH)


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def add_text(slide, x, y, w, h, text, size=20, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT_KR):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = rgb(color)
    run = paragraph.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run_properties = run._r.get_or_add_rPr()
    run_properties.set("lang", "ko-KR")
    for tag in ("a:ea", "a:cs"):
        element = run_properties.find(qn(tag))
        if element is None:
            element = etree.SubElement(run_properties, qn(tag))
        element.set("typeface", font)
    return box


def add_panel(slide, x, y, w, h, fill=NAVY_2, line=NAVY_3, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(1)
    return shape


def add_slide_chrome(slide, section: str, number: int, title: str, subtitle: str) -> None:
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(NAVY)
    add_text(slide, 0.58, 0.28, 3.2, 0.25, section.upper(), 9, MINT, True)
    add_text(slide, 0.58, 0.68, 12.0, 0.54, title, 27, WHITE, True)
    add_text(slide, 0.60, 1.24, 12.0, 0.32, subtitle, 11, PALE, False)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(1.63), Inches(12.16), Inches(0.025))
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(NAVY_3)
    line.line.fill.background()
    add_text(slide, 0.58, 7.20, 7.0, 0.18, "Source: model comparison portfolio workbook", 8, SLATE)
    add_text(slide, 12.35, 7.18, 0.38, 0.20, f"0{number}", 9, MINT, True, PP_ALIGN.RIGHT)


def add_label_value(slide, x, y, w, label, value, accent=MINT):
    add_text(slide, x, y, w, 0.24, label, 9, PALE, True)
    add_text(slide, x, y + 0.28, w, 0.46, value, 20, accent, True)


def krw_eok(value: float, decimals: int = 2, signed: bool = False) -> str:
    number = value / 100_000_000
    sign = "+" if signed and number > 0 else ""
    return f"{sign}₩{number:,.{decimals}f}억"


def build_presentation(metrics: dict) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.creator = "OpenAI Codex"
    prs.core_properties.title = "Sol High Investment Committee"
    prs.core_properties.subject = "GenOffice model comparison portfolio"
    blank = prs.slide_layouts[6]

    # Slide 1
    slide = prs.slides.add_slide(blank)
    add_slide_chrome(
        slide, "Executive decision frame", 1,
        "Base는 흑자, Downside는 자본 규율을 요구한다",
        "10개 투자안 전체 집행보다 Top 3 단계 투자가 위험 대비 효율적입니다.",
    )
    add_panel(slide, 0.58, 1.92, 12.16, 1.45)
    kpi_width = 2.20
    labels = [
        ("총 투자액", krw_eok(metrics["total_investment"])),
        ("예상 ARR", krw_eok(metrics["total_expected_arr"])),
        ("Base 순가치", krw_eok(metrics["base_net_value"], signed=True)),
        ("가중 성공확률", f'{metrics["weighted_success_probability"]:.1%}'),
        ("가중 이탈률", f'{metrics["weighted_attrition_probability"]:.1%}'),
    ]
    for idx, (label, value) in enumerate(labels):
        add_label_value(slide, 0.88 + idx * 2.35, 2.24, kpi_width, label, value)
    add_panel(slide, 0.58, 3.62, 7.80, 2.88, fill="0F1B36")
    add_text(slide, 0.88, 3.91, 2.6, 0.25, "DECISION", 10, MINT, True)
    add_text(slide, 0.88, 4.30, 6.9, 0.55, "전액 동시 집행 보류\nTop 3를 90일 단계 투자로 전환", 23, WHITE, True)
    add_text(
        slide, 0.88, 5.30, 6.95, 0.82,
        f'Downside {krw_eok(metrics["scenario_totals"]["Downside"], signed=True)} · '
        f'Base {krw_eok(metrics["scenario_totals"]["Base"], signed=True)} · '
        f'Upside {krw_eok(metrics["scenario_totals"]["Upside"], signed=True)}',
        14, PALE, False,
    )
    add_panel(slide, 8.66, 3.62, 4.08, 2.88, fill="12302F", line=MINT_DARK)
    add_text(slide, 8.98, 3.91, 2.8, 0.25, "WHY NOW", 10, MINT, True)
    add_text(slide, 8.98, 4.31, 3.35, 0.78, "Base 완충폭은\n총 투자액의 2.1%", 21, WHITE, True)
    add_text(slide, 8.98, 5.42, 3.20, 0.72, "성공확률이 흔들리면\nDownside 손실이 즉시 확대", 13, PALE)

    # Slide 2
    slide = prs.slides.add_slide(blank)
    add_slide_chrome(
        slide, "Portfolio & scenario stress test", 2,
        "Top 3가 Base 가치의 대부분을 만든다",
        "우선순위는 리스크조정점수, 시나리오는 동일 ARR·확률·비용 가정으로 계산했습니다.",
    )
    add_panel(slide, 0.58, 1.92, 5.08, 4.88, fill="0F1B36")
    add_text(slide, 0.88, 2.18, 3.6, 0.25, "TOP 3 PRIORITIES", 10, MINT, True)
    for idx, item in enumerate(metrics["top3"], start=1):
        y = 2.62 + (idx - 1) * 1.16
        add_text(slide, 0.88, y, 0.42, 0.36, f"{idx}", 18, MINT, True)
        add_text(slide, 1.42, y - 0.02, 3.70, 0.48, item["name"], 14, WHITE, True)
        add_text(
            slide, 1.42, y + 0.48, 3.75, 0.32,
            f'점수 {item["score"]:.1f}  ·  Base {krw_eok(item["base_net"], signed=True)}',
            10, PALE,
        )
    add_text(
        slide, 0.88, 6.26, 4.25, 0.30,
        f'Top 3 Base 합계  {krw_eok(metrics["top3_base_net"], signed=True)}',
        12, MINT, True,
    )

    add_panel(slide, 5.92, 1.92, 6.82, 3.25, fill="0F1B36")
    add_text(slide, 6.22, 2.18, 3.5, 0.25, "포트폴리오 순가치 (억원)", 10, MINT, True)
    baseline = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(6.32), Inches(3.60), Inches(5.98), Inches(0.018)
    )
    baseline.fill.solid()
    baseline.fill.fore_color.rgb = rgb(SLATE)
    baseline.line.fill.background()
    scenario_points = [
        ("Downside", metrics["scenario_totals"]["Downside"] / 100_000_000, "F87171"),
        ("Base", metrics["scenario_totals"]["Base"] / 100_000_000, MINT),
        ("Upside", metrics["scenario_totals"]["Upside"] / 100_000_000, MINT),
    ]
    for idx, (name, value, color) in enumerate(scenario_points):
        x = 6.78 + idx * 1.78
        height = max(abs(value) * 0.15, 0.07)
        y = 3.60 - height if value >= 0 else 3.62
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.78), Inches(height)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(color)
        bar.line.fill.background()
        label_y = y - 0.28 if value >= 0 else y + height + 0.04
        add_text(slide, x - 0.18, label_y, 1.15, 0.24, f"{value:+.2f}", 9, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, x - 0.25, 4.66, 1.30, 0.24, name, 9, PALE, True, PP_ALIGN.CENTER)
    add_panel(slide, 5.92, 5.40, 6.82, 1.40, fill="12302F", line=MINT_DARK)
    add_text(slide, 6.22, 5.66, 2.8, 0.24, "STRESS-TEST READOUT", 10, MINT, True)
    add_text(
        slide, 6.22, 6.04, 5.95, 0.46,
        "Base 흑자는 얇고 Downside 손실은 큽니다.\n따라서 Top 3 외 자본은 검증 후 해제합니다.",
        14, WHITE, True,
    )

    # Slide 3
    slide = prs.slides.add_slide(blank)
    add_slide_chrome(
        slide, "Recommendation & 90-day gates", 3,
        "₩2.66억으로 시작하고, 증거가 쌓일 때만 다음 자본을 연다",
        "Top 3 투자 조합 · 최대 리스크 3개 · 30/60/90일 중단 기준",
    )
    add_panel(slide, 0.58, 1.92, 3.70, 4.88, fill="12302F", line=MINT_DARK)
    add_text(slide, 0.88, 2.18, 2.7, 0.25, "RECOMMENDED BASKET", 10, MINT, True)
    add_text(slide, 0.88, 2.60, 3.00, 0.80, "API 포털 + CRM 자동화\n+ HR 스킬 매칭", 18, WHITE, True)
    add_label_value(slide, 0.88, 3.75, 2.8, "12개월 투자액", krw_eok(metrics["top3_total_investment"]))
    add_label_value(slide, 0.88, 4.65, 2.8, "예상 ARR", krw_eok(metrics["top3_expected_arr"], 3))
    add_label_value(slide, 0.88, 5.55, 2.8, "Base 순가치", krw_eok(metrics["top3_base_net"], 3, signed=True))

    add_panel(slide, 4.55, 1.92, 3.75, 4.88, fill="0F1B36")
    add_text(slide, 4.85, 2.18, 2.8, 0.25, "LARGEST RISKS", 10, MINT, True)
    risk_items = sorted(metrics["ranked"], key=lambda item: item["risk"], reverse=True)[:3]
    for idx, item in enumerate(risk_items, start=1):
        y = 2.66 + (idx - 1) * 1.18
        add_text(slide, 4.85, y, 0.35, 0.30, f"0{idx}", 12, MINT, True)
        add_text(slide, 5.28, y - 0.02, 2.62, 0.43, item["name"], 13, WHITE, True)
        add_text(
            slide, 5.28, y + 0.46, 2.55, 0.28,
            f'리스크 {item["risk"]} · Base {krw_eok(item["base_net"], 3, signed=True)}',
            9, PALE,
        )

    add_panel(slide, 8.57, 1.92, 4.17, 4.88, fill="0F1B36")
    add_text(slide, 8.87, 2.18, 2.9, 0.25, "30 / 60 / 90 DAY GATES", 10, MINT, True)
    gates = [
        ("30D", "우선순위 잠금", "78% / 72% / 70%\n하락 또는 Top 3 이탈 시 중단"),
        ("60D", "비용·ARR 검증", "투자액 ≤ ₩2.66억\n예상 ARR ≥ ₩3.251억"),
        ("90D", "집행 해제", "목표 Base ₩0.591억\n0 이하이면 중단"),
    ]
    for idx, (day, label, detail) in enumerate(gates):
        y = 2.61 + idx * 1.18
        add_text(slide, 8.87, y, 0.60, 0.32, day, 12, MINT, True)
        add_text(slide, 9.54, y - 0.01, 2.65, 0.32, label, 13, WHITE, True)
        add_text(slide, 9.54, y + 0.38, 2.70, 0.50, detail, 9, PALE)
    add_text(slide, 8.87, 6.25, 3.25, 0.30, "STOP = 다음 자본 미집행", 11, MINT, True)

    assert len(prs.slides) == 3
    prs.save(PPTX_PATH)


PIL_FONT_PATH = "/System/Library/Fonts/AppleSDGothicNeo.ttc"


def pil_font(size: int, bold: bool = False):
    return ImageFont.truetype(PIL_FONT_PATH, size=size, index=6 if bold else 0)


def pil_text(draw, xy, text, size, color, bold=False, anchor=None, align="left", spacing=6):
    draw.multiline_text(
        xy, text, font=pil_font(size, bold), fill=f"#{color}",
        anchor=anchor, align=align, spacing=spacing,
    )


def pil_panel(draw, box, fill=NAVY_2, outline=NAVY_3, radius=28, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=f"#{fill}", outline=f"#{outline}", width=width)


def pil_chrome(draw, section, title, subtitle, number):
    pil_text(draw, (72, 42), section.upper(), 18, MINT, True)
    pil_text(draw, (72, 92), title, 42, WHITE, True)
    pil_text(draw, (74, 151), subtitle, 20, PALE)
    draw.rectangle((72, 196, 1528, 199), fill=f"#{NAVY_3}")
    pil_text(draw, (72, 868), "Source: model comparison portfolio workbook", 13, SLATE)
    pil_text(draw, (1528, 866), f"0{number}", 15, MINT, True, anchor="ra")


def make_slide_images(metrics: dict) -> list[Image.Image]:
    slides = []

    # Slide 1
    image = Image.new("RGB", (1600, 900), f"#{NAVY}")
    draw = ImageDraw.Draw(image)
    pil_chrome(
        draw, "Executive decision frame",
        "Base는 흑자, Downside는 자본 규율을 요구한다",
        "10개 투자안 전체 집행보다 Top 3 단계 투자가 위험 대비 효율적입니다.", 1,
    )
    pil_panel(draw, (72, 225, 1528, 390))
    kpis = [
        ("총 투자액", krw_eok(metrics["total_investment"])),
        ("예상 ARR", krw_eok(metrics["total_expected_arr"])),
        ("Base 순가치", krw_eok(metrics["base_net_value"], signed=True)),
        ("가중 성공확률", f'{metrics["weighted_success_probability"]:.1%}'),
        ("가중 이탈률", f'{metrics["weighted_attrition_probability"]:.1%}'),
    ]
    for idx, (label, value) in enumerate(kpis):
        cx = 210 + idx * 291
        pil_text(draw, (cx, 262), label, 17, PALE, True, anchor="ma")
        pil_text(draw, (cx, 315), value, 29, MINT, True, anchor="ma")
    pil_panel(draw, (72, 425, 990, 790), fill="0F1B36")
    pil_text(draw, (112, 465), "DECISION", 17, MINT, True)
    pil_text(draw, (112, 525), "전액 동시 집행 보류\nTop 3를 90일 단계 투자로 전환", 34, WHITE, True, spacing=14)
    pil_text(
        draw, (112, 665),
        f'Downside {krw_eok(metrics["scenario_totals"]["Downside"], signed=True)}   ·   '
        f'Base {krw_eok(metrics["scenario_totals"]["Base"], signed=True)}   ·   '
        f'Upside {krw_eok(metrics["scenario_totals"]["Upside"], signed=True)}',
        21, PALE,
    )
    pil_panel(draw, (1092, 425, 1528, 790), fill="12302F", outline=MINT_DARK)
    pil_text(draw, (1132, 465), "WHY NOW", 17, MINT, True)
    pil_text(draw, (1132, 530), "Base 완충폭은\n총 투자액의 2.1%", 31, WHITE, True, spacing=12)
    pil_text(draw, (1132, 680), "성공확률이 흔들리면\nDownside 손실이 즉시 확대", 20, PALE, spacing=8)
    slides.append(image)

    # Slide 2
    image = Image.new("RGB", (1600, 900), f"#{NAVY}")
    draw = ImageDraw.Draw(image)
    pil_chrome(
        draw, "Portfolio & scenario stress test",
        "Top 3가 Base 가치의 대부분을 만든다",
        "리스크조정점수 우선순위와 동일 ARR·확률·비용 가정의 스트레스 테스트", 2,
    )
    pil_panel(draw, (72, 225, 660, 820), fill="0F1B36")
    pil_text(draw, (112, 265), "TOP 3 PRIORITIES", 17, MINT, True)
    for idx, item in enumerate(metrics["top3"], start=1):
        y = 335 + (idx - 1) * 135
        pil_text(draw, (112, y), str(idx), 31, MINT, True)
        pil_text(draw, (165, y), item["name"], 25, WHITE, True)
        pil_text(
            draw, (165, y + 50),
            f'점수 {item["score"]:.1f}  ·  Base {krw_eok(item["base_net"], signed=True)}',
            17, PALE,
        )
    pil_text(draw, (112, 760), f'Top 3 Base 합계  {krw_eok(metrics["top3_base_net"], signed=True)}', 20, MINT, True)
    pil_panel(draw, (700, 225, 1528, 585), fill="0F1B36")
    pil_text(draw, (740, 265), "포트폴리오 순가치 (억원)", 17, MINT, True)
    baseline_y = 410
    draw.rectangle((770, baseline_y, 1475, baseline_y + 2), fill=f"#{SLATE}")
    scenario_points = [
        ("Downside", metrics["scenario_totals"]["Downside"] / 100_000_000, "F87171"),
        ("Base", metrics["scenario_totals"]["Base"] / 100_000_000, MINT),
        ("Upside", metrics["scenario_totals"]["Upside"] / 100_000_000, MINT),
    ]
    for idx, (name, value, color) in enumerate(scenario_points):
        x = 850 + idx * 245
        height = max(int(abs(value) * 20), 8)
        top = baseline_y - height if value >= 0 else baseline_y + 2
        bottom = baseline_y if value >= 0 else baseline_y + 2 + height
        draw.rounded_rectangle((x, top, x + 90, bottom), radius=8, fill=f"#{color}")
        value_y = top - 12 if value >= 0 else bottom - 28
        pil_text(draw, (x + 45, value_y), f"{value:+.2f}", 17, WHITE, True, anchor="ms")
        pil_text(draw, (x + 45, 565), name, 17, PALE, True, anchor="ms")
    pil_panel(draw, (700, 615, 1528, 820), fill="12302F", outline=MINT_DARK)
    pil_text(draw, (740, 650), "STRESS-TEST READOUT", 17, MINT, True)
    pil_text(
        draw, (740, 700),
        "Base 흑자는 얇고 Downside 손실은 큽니다.\n따라서 Top 3 외 자본은 검증 후 해제합니다.",
        26, WHITE, True, spacing=10,
    )
    slides.append(image)

    # Slide 3
    image = Image.new("RGB", (1600, 900), f"#{NAVY}")
    draw = ImageDraw.Draw(image)
    pil_chrome(
        draw, "Recommendation & 90-day gates",
        "₩2.66억으로 시작하고, 증거가 쌓일 때만 다음 자본을 연다",
        "Top 3 투자 조합 · 최대 리스크 3개 · 30/60/90일 중단 기준", 3,
    )
    pil_panel(draw, (72, 225, 535, 820), fill="12302F", outline=MINT_DARK)
    pil_text(draw, (112, 265), "RECOMMENDED BASKET", 17, MINT, True)
    pil_text(draw, (112, 325), "API 포털 + CRM 자동화\n+ HR 스킬 매칭", 29, WHITE, True, spacing=10)
    basket_metrics = [
        ("12개월 투자액", krw_eok(metrics["top3_total_investment"])),
        ("예상 ARR", krw_eok(metrics["top3_expected_arr"], 3)),
        ("Base 순가치", krw_eok(metrics["top3_base_net"], 3, signed=True)),
    ]
    for idx, (label, value) in enumerate(basket_metrics):
        y = 485 + idx * 102
        pil_text(draw, (112, y), label, 16, PALE, True)
        pil_text(draw, (112, y + 35), value, 30, MINT, True)

    pil_panel(draw, (565, 225, 1045, 820), fill="0F1B36")
    pil_text(draw, (605, 265), "LARGEST RISKS", 17, MINT, True)
    risk_items = sorted(metrics["ranked"], key=lambda item: item["risk"], reverse=True)[:3]
    for idx, item in enumerate(risk_items, start=1):
        y = 340 + (idx - 1) * 145
        pil_text(draw, (605, y), f"0{idx}", 22, MINT, True)
        pil_text(draw, (665, y), item["name"], 22, WHITE, True)
        pil_text(
            draw, (665, y + 48),
            f'리스크 {item["risk"]}  ·  Base {krw_eok(item["base_net"], 3, signed=True)}',
            16, PALE,
        )

    pil_panel(draw, (1075, 225, 1528, 820), fill="0F1B36")
    pil_text(draw, (1115, 265), "30 / 60 / 90 DAY GATES", 17, MINT, True)
    gates = [
        ("30D", "우선순위 잠금", "78% / 72% / 70%\nTop 3 이탈 시 중단"),
        ("60D", "비용·ARR 검증", "투자액 ≤ ₩2.66억\n예상 ARR ≥ ₩3.251억"),
        ("90D", "집행 해제", "목표 Base ₩0.591억\n0 이하이면 중단"),
    ]
    for idx, (day, label, detail) in enumerate(gates):
        y = 340 + idx * 145
        pil_text(draw, (1115, y), day, 24, MINT, True)
        pil_text(draw, (1190, y), label, 20, WHITE, True)
        pil_text(draw, (1190, y + 42), detail, 16, PALE, spacing=5)
    pil_text(draw, (1115, 760), "STOP = 다음 자본 미집행", 18, MINT, True)
    slides.append(image)
    return slides


def build_raster_presentation(metrics: dict) -> None:
    slides = make_slide_images(metrics)
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.creator = "OpenAI Codex"
    prs.core_properties.title = "Sol High Investment Committee"
    prs.core_properties.subject = "GenOffice model comparison portfolio"
    blank = prs.slide_layouts[6]
    for index, image in enumerate(slides, start=1):
        image.save(OUT_DIR / f"sol-high-slide-{index}.png", format="PNG", optimize=True)
        buffer = BytesIO()
        image.save(buffer, format="PNG", optimize=True)
        buffer.seek(0)
        slide = prs.slides.add_slide(blank)
        picture = slide.shapes.add_picture(buffer, 0, 0, width=prs.slide_width, height=prs.slide_height)
        picture.name = f"Slide {index} rendered content"
    assert len(prs.slides) == 3
    prs.save(PPTX_PATH)


def render_dashboard_png(metrics: dict) -> None:
    image = Image.new("RGB", (1600, 1000), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((32, 30, 1568, 110), radius=16, fill=f"#{NAVY}")
    pil_text(draw, (64, 58), "B2B SaaS 투자 포트폴리오 대시보드", 34, WHITE, True)
    kpis = [
        ("총 투자액", krw_eok(metrics["total_investment"])),
        ("예상 ARR", krw_eok(metrics["total_expected_arr"])),
        ("Base 순가치", krw_eok(metrics["base_net_value"], signed=True)),
        ("가중 성공확률", f'{metrics["weighted_success_probability"]:.1%}'),
        ("가중 이탈률", f'{metrics["weighted_attrition_probability"]:.1%}'),
    ]
    for idx, (label, value) in enumerate(kpis):
        x1 = 32 + idx * 307
        x2 = x1 + 285
        draw.rounded_rectangle((x1, 135, x2, 235), radius=15, fill="#ECFDF5", outline=f"#{MINT_DARK}", width=2)
        pil_text(draw, ((x1 + x2) // 2, 160), label, 16, NAVY, True, anchor="ma")
        pil_text(draw, ((x1 + x2) // 2, 206), value, 24, MINT_DARK, True, anchor="ma")
    pil_panel(draw, (32, 265, 780, 745), fill="FFFFFF", outline="CBD5E1", radius=18)
    pil_text(draw, (66, 300), "우선순위 Top 5", 23, NAVY, True)
    for idx, item in enumerate(metrics["ranked"][:5], start=1):
        y = 370 + (idx - 1) * 68
        pil_text(draw, (68, y), str(idx), 19, MINT_DARK, True)
        pil_text(draw, (112, y), item["name"], 18, NAVY, True)
        pil_text(draw, (585, y), f'{item["score"]:.1f}', 18, NAVY, True, anchor="ra")
        bar_width = int(item["score"] / 70 * 65)
        draw.rounded_rectangle((610, y + 3, 610 + bar_width, y + 23), radius=8, fill=f"#{MINT_DARK}")
        pil_text(draw, (750, y), krw_eok(item["base_net"], 3, signed=True), 14, NAVY, anchor="ra")
        draw.line((66, y + 46, 744, y + 46), fill="#E2E8F0", width=1)
    pil_panel(draw, (810, 265, 1568, 745), fill="FFFFFF", outline="CBD5E1", radius=18)
    pil_text(draw, (844, 300), "시나리오 스트레스 테스트", 23, NAVY, True)
    baseline_y = 520
    draw.line((875, baseline_y, 1500, baseline_y), fill=f"#{SLATE}", width=2)
    scenario_points = [
        ("Downside", metrics["scenario_totals"]["Downside"] / 100_000_000, "F87171"),
        ("Base", metrics["scenario_totals"]["Base"] / 100_000_000, MINT),
        ("Upside", metrics["scenario_totals"]["Upside"] / 100_000_000, MINT),
    ]
    for idx, (name, value, color) in enumerate(scenario_points):
        x = 930 + idx * 215
        height = max(int(abs(value) * 29), 8)
        top = baseline_y - height if value >= 0 else baseline_y + 2
        bottom = baseline_y if value >= 0 else baseline_y + 2 + height
        draw.rounded_rectangle((x, top, x + 90, bottom), radius=7, fill=f"#{color}")
        value_y = top - 12 if value >= 0 else bottom - 25
        value_color = NAVY if value >= 0 else WHITE
        pil_text(draw, (x + 45, value_y), f"{value:+.2f}억", 17, value_color, True, anchor="ms")
        pil_text(draw, (x + 45, 715), name, 17, NAVY, True, anchor="ms")
    draw.rounded_rectangle((32, 780, 1568, 945), radius=18, fill="#ECFDF5", outline=f"#{MINT_DARK}", width=2)
    pil_text(draw, (66, 815), "경영진 권고", 20, MINT_DARK, True)
    pil_text(
        draw, (66, 860),
        "Downside에서 순가치가 음수이므로 전액 동시 집행보다 1~3위 투자안의 단계적 집행을 권고합니다.",
        28, NAVY, True,
    )
    image.save(OUT_DIR / "sol-high-dashboard.png", format="PNG", optimize=True)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    items, assumptions = extract_source()
    metrics = compute_metrics(items, assumptions)
    build_workbook(items, assumptions, metrics)
    build_raster_presentation(metrics)
    render_dashboard_png(metrics)
    serializable = {
        key: value for key, value in metrics.items()
        if key not in {"ranked", "top3", "scenario_values"}
    }
    serializable["top3"] = [
        {
            "rank": item["rank"],
            "name": item["name"],
            "score": item["score"],
            "base_net": item["base_net"],
            "success_probability": item["prob"],
        }
        for item in metrics["top3"]
    ]
    serializable["largest_risks"] = [
        {"name": item["name"], "risk_score": item["risk"], "base_net": item["base_net"]}
        for item in sorted(metrics["ranked"], key=lambda item: item["risk"], reverse=True)[:3]
    ]
    METRICS_PATH.write_text(json.dumps(serializable, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(serializable, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
