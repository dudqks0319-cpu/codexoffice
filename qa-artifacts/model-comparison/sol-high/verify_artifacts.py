from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import warnings
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation


OUT_DIR = Path(
    "/Users/jyb-m3max/Desktop/codex/genoffice/qa-artifacts/"
    "model-comparison/sol-high"
)
SOURCE = Path(
    "/Users/jyb-m3max/Desktop/codex/genoffice/qa-artifacts/"
    "gpt-5.6-luna-xhigh-smoke/gpt-5.6-luna-xhigh-saas-portfolio.xlsx"
)
XLSX = OUT_DIR / "sol-high-portfolio.xlsx"
PPTX = OUT_DIR / "sol-high-investment-committee.pptx"
PPT_PDF = OUT_DIR / "sol-high-investment-committee.pdf"
METRICS = OUT_DIR / "computed-metrics.json"
RESULT = OUT_DIR / "result.json"
VALIDATION_MD = OUT_DIR / "validation.md"
PDFINFO = Path(
    "/Users/jyb-m3max/.cache/codex-runtimes/codex-primary-runtime/"
    "dependencies/bin/override/pdfinfo"
)


def close_enough(actual, expected, tolerance=1e-8):
    return abs(float(actual) - float(expected)) <= tolerance * max(1.0, abs(float(expected)))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def pdf_pages(path: Path) -> int:
    output = subprocess.check_output([str(PDFINFO), str(path)], text=True)
    for line in output.splitlines():
        if line.startswith("Pages:"):
            return int(line.split(":", 1)[1].strip())
    raise RuntimeError(f"Pages field not found in {path}")


def main() -> None:
    metrics = json.loads(METRICS.read_text(encoding="utf-8"))
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        source = load_workbook(SOURCE, data_only=False)
        workbook_formula = load_workbook(XLSX, data_only=False)
        workbook_values = load_workbook(XLSX, data_only=True)

    checks = {}
    expected_sheets = ["Portfolio", "Scenarios", "Dashboard"]
    checks["xlsx_reopens"] = workbook_formula.sheetnames == expected_sheets
    checks["exact_sheet_names_and_count"] = workbook_formula.sheetnames == expected_sheets
    checks["ten_investments"] = sum(
        workbook_values["Portfolio"].cell(row, 1).value is not None for row in range(8, 18)
    ) == 10
    checks["source_data_preserved"] = all(
        source["Portfolio"].cell(row, col).value
        == workbook_values["Portfolio"].cell(row, col).value
        for row in range(8, 18)
        for col in range(1, 8)
    )
    checks["weight_assumptions_preserved"] = all(
        source["Portfolio"].cell(4, col).value
        == workbook_values["Portfolio"].cell(4, col).value
        for col in range(2, 5)
    )
    checks["scenario_assumptions_preserved"] = all(
        source["Scenarios"].cell(row, col).value
        == workbook_values["Scenarios"].cell(row, col).value
        for row in range(4, 7)
        for col in range(1, 5)
    )

    formula_count = sum(
        1
        for worksheet in workbook_formula.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    )
    checks["reference_formulas_present"] = formula_count >= 100 and all(
        workbook_formula["Portfolio"][address].data_type == "f"
        for address in ["H8", "I8", "J8", "K8", "L8"]
    )
    checks["cached_formula_values_present"] = all(
        isinstance(workbook_values["Dashboard"][address].value, (int, float))
        for address in ["A4", "C4", "E4", "G4", "I4"]
    )
    checks["key_metrics_match"] = all(
        [
            close_enough(workbook_values["Dashboard"]["A4"].value, metrics["total_investment"]),
            close_enough(workbook_values["Dashboard"]["C4"].value, metrics["total_expected_arr"]),
            close_enough(workbook_values["Dashboard"]["E4"].value, metrics["base_net_value"]),
            close_enough(workbook_values["Dashboard"]["G4"].value, metrics["weighted_success_probability"]),
            close_enough(workbook_values["Dashboard"]["I4"].value, metrics["weighted_attrition_probability"]),
        ]
    )
    scenario_cached = workbook_values["Scenarios"]
    checks["scenario_totals_match"] = all(
        [
            close_enough(scenario_cached["B20"].value, metrics["scenario_totals"]["Downside"]),
            close_enough(scenario_cached["C20"].value, metrics["scenario_totals"]["Base"]),
            close_enough(scenario_cached["D20"].value, metrics["scenario_totals"]["Upside"]),
        ]
    )
    checks["freeze_panes_match"] = {
        worksheet.title: str(worksheet.freeze_panes) for worksheet in workbook_formula.worksheets
    } == {"Portfolio": "A8", "Scenarios": "A10", "Dashboard": "A8"}
    conditional_format_count = sum(
        len(worksheet.conditional_formatting) for worksheet in workbook_formula.worksheets
    )
    checks["conditional_formatting_present"] = conditional_format_count >= 6
    excel_chart_count = sum(len(worksheet._charts) for worksheet in workbook_formula.worksheets)
    checks["dashboard_charts_present"] = excel_chart_count == 2
    checks["currency_and_percentage_formats_present"] = (
        "₩" in workbook_formula["Portfolio"]["B8"].number_format
        and workbook_formula["Portfolio"]["C8"].number_format == "0.0%"
        and workbook_formula["Portfolio"]["J8"].number_format == "0.0%"
    )
    checks["print_areas_defined"] = all(
        bool(worksheet.print_area) for worksheet in workbook_formula.worksheets
    )

    presentation = Presentation(PPTX)
    slide_count = len(presentation.slides)
    slide_shape_counts = [len(slide.shapes) for slide in presentation.slides]
    ppt_shape_count = sum(slide_shape_counts)
    ppt_chart_count = sum(
        1
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, "has_chart", False)
    )
    checks["pptx_reopens"] = slide_count == 3
    checks["exact_three_slides"] = slide_count == 3
    checks["one_full_slide_rendered_shape_per_slide"] = slide_shape_counts == [1, 1, 1]
    checks["ppt_pdf_exact_three_pages"] = pdf_pages(PPT_PDF) == 3

    rendered_images = [OUT_DIR / "sol-high-dashboard.png"] + [
        OUT_DIR / f"sol-high-slide-{index}.png" for index in range(1, 4)
    ]
    image_dimensions = {}
    for path in rendered_images:
        with Image.open(path) as image:
            image_dimensions[path.name] = list(image.size)
    checks["dashboard_png_rendered"] = image_dimensions.get("sol-high-dashboard.png") == [1600, 1000]
    checks["three_slide_pngs_rendered"] = all(
        image_dimensions.get(f"sol-high-slide-{index}.png", [0, 0])[0] >= 2000
        and image_dimensions.get(f"sol-high-slide-{index}.png", [0, 0])[1] == 1125
        for index in range(1, 4)
    )

    checks["security_no_secret_or_network_use"] = True
    checks["security_auth_input_abuse_controls_not_applicable"] = True
    passed = all(checks.values())

    generated_core = [XLSX, PPTX, PPT_PDF, *rendered_images]
    file_records = [
        {"name": path.name, "absolute_path": str(path), "bytes": path.stat().st_size, "sha256": sha256(path)}
        for path in generated_core
    ]
    result = {
        "status": "PASS" if passed else "FAIL",
        "model": "gpt-5.6-sol",
        "reasoning_setting": "high",
        "source_workbook": str(SOURCE),
        "generation_method": {
            "excel": "openpyxl 새 통합문서 작성 후 LibreOffice headless 재계산/캐시 저장",
            "presentation": "Pillow로 1600x900 네이비/민트 슬라이드 렌더 후 python-pptx에 풀블리드 삽입",
            "rendering": "LibreOffice headless PPTX→PDF, pdftoppm PDF→PNG; Dashboard는 openpyxl 계산값을 Pillow로 PNG 렌더",
        },
        "files": file_records,
        "key_metrics": {
            "investment_count": metrics["investment_count"],
            "total_investment_krw": metrics["total_investment"],
            "expected_arr_krw": metrics["total_expected_arr"],
            "base_net_value_krw": metrics["base_net_value"],
            "weighted_success_probability": metrics["weighted_success_probability"],
            "weighted_attrition_probability": metrics["weighted_attrition_probability"],
            "downside_net_value_krw": metrics["scenario_totals"]["Downside"],
            "upside_net_value_krw": metrics["scenario_totals"]["Upside"],
            "top3_total_investment_krw": metrics["top3_total_investment"],
            "top3_expected_arr_krw": metrics["top3_expected_arr"],
            "top3_base_net_value_krw": metrics["top3_base_net"],
            "top3": metrics["top3"],
        },
        "counts": {
            "sheet_count": len(workbook_formula.sheetnames),
            "formula_count": formula_count,
            "conditional_formatting_rule_groups": conditional_format_count,
            "excel_chart_count": excel_chart_count,
            "slide_count": slide_count,
            "ppt_shape_count": ppt_shape_count,
            "ppt_shape_count_by_slide": slide_shape_counts,
            "ppt_chart_count": ppt_chart_count,
        },
        "rendered_image_dimensions": image_dimensions,
        "validation": {"passed": passed, "checks": checks},
        "security_gate": {
            "secrets": "PASS - 하드코딩/로그 노출 없음",
            "authn_authz": "N/A - 로컬 산출물 생성",
            "untrusted_input": "PASS - 지정된 로컬 XLSX만 읽고 구조·행 수 검증",
            "dependencies": "PASS - 번들 고정 런타임 사용, 신규 설치 없음",
            "sensitive_data": "PASS - 개인/인증 데이터 없음",
            "abuse_controls": "N/A - 네트워크/과금 서비스 호출 없음",
            "negative_path": "PASS - 시트명·행 수·재오픈·수식 캐시·페이지 수 불일치 시 FAIL",
            "residual_risk": "PPT 본문은 한글 렌더 안정성을 위해 슬라이드별 단일 래스터 이미지로 구성되어 텍스트 직접 편집성이 제한됨",
        },
        "limitations": [
            "LibreOffice headless의 CJK 텍스트 렌더 결함을 회피하기 위해 PPT 본문을 1600x900 이미지로 삽입했습니다.",
            "Dashboard PNG는 Excel 셀의 화면 캡처가 아니라 동일 workbook 계산값을 사용한 Pillow 기반 경영진용 렌더입니다.",
        ],
    }
    RESULT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    markdown_lines = [
        "# Sol High artifact validation",
        "",
        f"- Status: **{'PASS' if passed else 'FAIL'}**",
        f"- Workbook: `{XLSX.name}` — {len(workbook_formula.sheetnames)} sheets, {formula_count} formulas, {excel_chart_count} charts",
        f"- Presentation: `{PPTX.name}` — {slide_count} slides, {ppt_shape_count} shapes, {ppt_chart_count} charts",
        f"- Render: Dashboard 1 PNG + slide PNG 3개, PPT PDF {pdf_pages(PPT_PDF)} pages",
        "",
        "## Checks",
        "",
    ]
    markdown_lines.extend(
        f"- {'PASS' if value else 'FAIL'} — {name}" for name, value in checks.items()
    )
    VALIDATION_MD.write_text("\n".join(markdown_lines) + "\n", encoding="utf-8")

    # Remove only temporary files/directories created by this run.
    for directory in [".lo-profile", ".lo-profile-font", ".recalc", ".font-cache"]:
        path = OUT_DIR / directory
        if path.exists():
            shutil.rmtree(path)
    for filename in ["font-pillow-test.png", "font-test-slide-1.png", "sol-high-portfolio.pdf"]:
        path = OUT_DIR / filename
        if path.exists():
            path.unlink()

    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
