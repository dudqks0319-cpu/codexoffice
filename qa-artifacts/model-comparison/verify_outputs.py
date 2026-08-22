#!/usr/bin/env python3
import json
import math
import sys
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation


EXPECTED_SHEETS = ["Portfolio", "Scenarios", "Dashboard"]


def workbook_metrics(path: Path) -> dict:
    wb = load_workbook(path, data_only=False)
    formulas = []
    cross_sheet = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append(f"{ws.title}!{cell.coordinate}:{cell.value}")
                    if "!" in cell.value:
                        cross_sheet.append(f"{ws.title}!{cell.coordinate}:{cell.value}")

    ws = wb["Portfolio"]
    rows = []
    for row in range(8, 18):
        values = [ws.cell(row, col).value for col in range(1, 8)]
        rows.append(values)

    total_investment = sum(float(row[3]) + 12 * float(row[4]) for row in rows)
    expected_arr = sum(float(row[1]) * float(row[2]) for row in rows)
    base_value = expected_arr - total_investment
    scenarios = {}
    for name, arr_mult, probability_delta, cost_mult in (
        ("Downside", 0.8, -0.15, 1.15),
        ("Base", 1.0, 0.0, 1.0),
        ("Upside", 1.2, 0.1, 0.9),
    ):
        scenarios[name] = sum(
            float(row[1]) * arr_mult * min(1.0, max(0.0, float(row[2]) + probability_delta))
            - (float(row[3]) + 12 * float(row[4])) * cost_mult
            for row in rows
        )

    freeze_panes = {ws.title: str(ws.freeze_panes) if ws.freeze_panes else None for ws in wb.worksheets}
    has_freeze_panes = all(value is not None for value in freeze_panes.values())
    return {
        "path": str(path),
        "sheets": wb.sheetnames,
        "sheet_count": len(wb.sheetnames),
        "portfolio_rows": len(rows),
        "portfolio_input_fingerprint": rows,
        "formula_count": len(formulas),
        "cross_sheet_formula_count": len(cross_sheet),
        "formula_samples": formulas[:12],
        "cross_sheet_samples": cross_sheet[:12],
        "total_investment": total_investment,
        "expected_arr": expected_arr,
        "base_value": base_value,
        "scenario_values": scenarios,
        "has_conditional_formatting": any(
            len(ws.conditional_formatting) > 0 for ws in wb.worksheets
        ),
        "freeze_panes": freeze_panes,
        "has_freeze_panes": has_freeze_panes,
        "valid": wb.sheetnames == EXPECTED_SHEETS
        and len(rows) == 10
        and len(formulas) >= 50
        and len(cross_sheet) >= 20
        and has_freeze_panes,
    }


def presentation_metrics(path: Path) -> dict:
    prs = Presentation(path)
    slide_text = []
    shapes = 0
    charts = 0
    pictures = 0
    shapes_per_slide = []
    for index, slide in enumerate(prs.slides, start=1):
        texts = []
        slide_shapes = 0
        for shape in slide.shapes:
            shapes += 1
            slide_shapes += 1
            if getattr(shape, "has_chart", False):
                charts += 1
            if shape.shape_type == 13:
                pictures += 1
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        slide_text.append({"slide": index, "text": texts})
        shapes_per_slide.append(slide_shapes)
    return {
        "path": str(path),
        "slide_count": len(prs.slides),
        "shape_count": shapes,
        "chart_count": charts,
        "picture_count": pictures,
        "shapes_per_slide": shapes_per_slide,
        "image_only": pictures == len(prs.slides) and shapes == len(prs.slides),
        "slide_text": slide_text,
        "valid": len(prs.slides) == 3 and all(count >= 1 for count in shapes_per_slide),
    }


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit("usage: verify_outputs.py sol.xlsx sol.pptx luna.xlsx luna.pptx")
    sol_xlsx, sol_pptx, luna_xlsx, luna_pptx = map(Path, sys.argv[1:])
    output = {
        "sol_high": {
            "workbook": workbook_metrics(sol_xlsx),
            "presentation": presentation_metrics(sol_pptx),
        },
        "luna_max": {
            "workbook": workbook_metrics(luna_xlsx),
            "presentation": presentation_metrics(luna_pptx),
        },
    }
    left = output["sol_high"]["workbook"]
    right = output["luna_max"]["workbook"]
    output["comparison"] = {
        "same_portfolio_inputs": left["portfolio_input_fingerprint"]
        == right["portfolio_input_fingerprint"],
        "same_computed_metrics": all(
            math.isclose(left[key], right[key], rel_tol=0, abs_tol=0.01)
            for key in ("total_investment", "expected_arr", "base_value")
        ),
        "both_valid": all(
            output[name][kind]["valid"]
            for name in ("sol_high", "luna_max")
            for kind in ("workbook", "presentation")
        ),
    }
    rendered = json.dumps(output, ensure_ascii=False, indent=2)
    result_path = Path(__file__).with_name("independent-verification.json")
    result_path.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)


if __name__ == "__main__":
    main()
